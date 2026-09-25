"""Deterministic DOCX report analyzer for GIP."""
from __future__ import annotations
from dataclasses import dataclass, field
import re
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from docx import Document
from .model import Defect, InspectionReport, ReportMetadata

_CONTRACT_RE = re.compile(r"(?:договор(?:а|у)?|контракт(?:а|у)?)[^\d№]{0,20}(?:№\s*)?([A-Za-zА-Яа-я0-9./_-]+)", re.I)
_DATE_RE = re.compile(r"\b(?:\d{1,2}[./-]\d{1,2}[./-]\d{2,4}|\d{4}-\d{2}-\d{2})\b")
_DEFECT_HINTS = ("трещин","трещина","разруш","корроз","отсутств","повреж","замач","вымыва","обруш","деформац","прогиб","скол","ослаб")

@dataclass(frozen=True)
class AnalysisResult:
    report: InspectionReport
    paragraphs: list[str] = field(default_factory=list)
    tables: list[list[list[str]]] = field(default_factory=list)
    candidate_defect_paragraphs: list[str] = field(default_factory=list)
    contract_candidates: list[str] = field(default_factory=list)
    address_candidates: list[str] = field(default_factory=list)
    date_candidates: list[str] = field(default_factory=list)
    document_part_text: list[tuple[str,str]] = field(default_factory=list)

class DocxReportAnalyzer:
    def analyze(self, path: str | Path) -> AnalysisResult:
        p=Path(path)
        if p.suffix.lower() != ".docx": raise ValueError("analyzer accepts .docx files only")
        doc=Document(p)
        paragraphs=[x.text.strip() for x in doc.paragraphs if x.text.strip()]
        tables=[[[c.text.strip() for c in row.cells] for row in t.rows] for t in doc.tables]
        parts=self._document_parts(p)
        all_text="\n".join(text for _,text in parts)
        contracts=_unique(m.group(1) for m in _CONTRACT_RE.finditer(all_text))
        dates=_unique(_DATE_RE.findall(all_text))
        addresses=_address_candidates(parts)
        address=self._address(paragraphs)
        candidates=[x for x in paragraphs if any(h in x.lower() for h in _DEFECT_HINTS)]
        defects=[Defect(id=None,structure="",description=x) for x in candidates]
        return AnalysisResult(
            InspectionReport(metadata=ReportMetadata(contract_number=contracts[0] if contracts else None,address=address,dates=dates),defects=defects),
            paragraphs,tables,candidates,contracts,addresses,dates,parts)

    @staticmethod
    def _document_parts(path: Path) -> list[tuple[str,str]]:
        parts=[]
        with ZipFile(path) as z:
            for name in z.namelist():
                if not name.startswith("word/") or not name.endswith(".xml"): continue
                try: root=ET.fromstring(z.read(name))
                except ET.ParseError: continue
                text=" ".join(t.text for t in root.iter() if t.tag.endswith("}t") and t.text).strip()
                if text: parts.append((name,text))
        return parts

    @staticmethod
    def _address(paragraphs: list[str]) -> str|None:
        for text in paragraphs:
            if "адрес" in text.lower() and ":" in text:
                return text.split(":",1)[1].strip()
        return None

def _unique(values) -> list[str]:
    out=[]; seen=set()
    for value in values:
        key=" ".join(re.findall(r"[a-zа-я0-9]+",(value or "").lower().replace("ё","е")))
        if key and key not in seen:
            seen.add(key); out.append(value)
    return out

def _address_candidates(parts: list[tuple[str,str]]) -> list[str]:
    values=[]
    pattern=re.compile(
        r"(?:адрес|место\s+расположения)\s*[:№-]?\s*"
        r"((?:Россия\s*,?\s*)?(?:[А-ЯЁа-яёA-Za-z-]+\s+область\s*,?\s*)?"
        r"(?:г\.?\s*)?[А-ЯЁа-яёA-Za-z-]+\s*,\s*"
        r"(?:ул\.?\s*)?[А-ЯЁа-яё0-9 .-]+,\s*\d+[А-Яа-яA-Za-z]?)",
        re.I)
    for _,text in parts:
        for m in pattern.finditer(text):
            value=m.group(1).strip(" .;,:»\"")
            if value: values.append(value)
    return _unique(values)
