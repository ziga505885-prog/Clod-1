"""Deterministic post-patch and document verification."""
from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
from docx import Document
from docx.shared import RGBColor
from .patches import Patch, PatchKind

@dataclass(frozen=True)
class VerificationFinding:
    code: str
    message: str

@dataclass(frozen=True)
class DocumentVerification:
    passed: bool
    findings: list[VerificationFinding]

class DocxVerificationEngine:
    def verify_patch(self, path: str | Path, patch: Patch) -> DocumentVerification:
        p = Path(path)
        if p.suffix.lower() != ".docx":
            raise ValueError("verification engine accepts .docx files only")
        doc = Document(p)
        findings = []
        locations = self._runs(doc)
        old_count = sum(r.text.count(patch.old) for r in locations)
        new_count = sum(r.text.count(patch.new) for r in locations)
        new_runs = [r for r in locations if patch.new in r.text]
        if old_count:
            findings.append(VerificationFinding("OLD_TEXT_PRESENT", "Original text is still present"))
        if new_count != 1:
            findings.append(VerificationFinding("REPLACEMENT_COUNT", f"Expected one replacement, found {new_count}"))
        elif not self._color_ok(new_runs[0], patch.kind):
            findings.append(VerificationFinding("COLOR_MISMATCH", "Replacement has incorrect GIP color semantics"))
        if patch.kind in (PatchKind.ADDRESS, PatchKind.DATE) and patch.reason:
            findings.append(VerificationFinding("SPECIAL_REASON", "Address/date patch must not contain explanatory text"))
        return DocumentVerification(not findings, findings)

    def verify_document(self, path: str | Path) -> DocumentVerification:
        p = Path(path)
        if p.suffix.lower() != ".docx":
            raise ValueError("verification engine accepts .docx files only")
        doc = Document(p)
        findings = []
        for i, table in enumerate(doc.tables):
            if not table.rows:
                continue
            headers = [c.text.strip().lower() for c in table.rows[0].cells]
            for j, header in enumerate(headers):
                if header == "категория состояния":
                    for ri, row in enumerate(table.rows[1:], 1):
                        value = row.cells[j].text.strip() if j < len(row.cells) else ""
                        if value not in {"аварийное", "работоспособное", "ограниченно-работоспособное", "нормативное"}:
                            findings.append(VerificationFinding("INVALID_CATEGORY", f"Table {i}, row {ri}: invalid category"))
        return DocumentVerification(not findings, findings)

    def _runs(self, doc):
        out = [r for p in doc.paragraphs for r in p.runs]
        out += [r for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs]
        return out

    @staticmethod
    def _color_ok(run, kind):
        rgb = run.font.color.rgb
        expected = RGBColor(0, 0, 255) if kind in (PatchKind.ADDRESS, PatchKind.DATE) else RGBColor(0, 128, 0)
        return rgb == expected
