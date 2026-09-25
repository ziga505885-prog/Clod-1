from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import re

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None


@dataclass(frozen=True)
class SectionSoilBinding:
    section: str
    lira_point: tuple[float, float] | None = None
    soil_point: tuple[float, float] | None = None
    azimuth: float | None = None
    height: float | None = None
    notes: tuple[str, ...] = ()


@dataclass(frozen=True)
class LoadCase:
    number: str
    name: str
    load_type: str
    duration_share: float | None
    reliability_factor: float | None


@dataclass
class CalculationWorkbookProfile:
    bindings: list[SectionSoilBinding] = field(default_factory=list)
    load_cases: list[LoadCase] = field(default_factory=list)


_FLOAT_RE = re.compile(r"^-?\d+(?:[.,]\d+)?$")


def _float(value):
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str) and _FLOAT_RE.fullmatch(value.strip()):
        return float(value.strip().replace(",", "."))
    return None


def _section_number(value: object) -> str | None:
    m = re.search(r"секция\s+(\d+)", str(value), re.I)
    return m.group(1) if m else None


def read_section_soil_bindings(path: str | Path) -> list[SectionSoilBinding]:
    if load_workbook is None:
        raise RuntimeError("openpyxl is required to read calculation workbooks")
    ws = load_workbook(path, data_only=True, read_only=True).active
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    result: list[SectionSoilBinding] = []
    current: dict | None = None
    for row in rows:
        blob = " ".join(str(x) for x in row if x is not None)
        sec = _section_number(blob)
        if sec:
            if current:
                result.append(SectionSoilBinding(**current))
            current = {"section": sec, "notes": []}
            continue
        if current is None:
            continue
        if row and str(row[0]).strip().lower() == "точка 1":
            current["lira_point"] = (_float(row[1]), _float(row[2]))
        elif row and str(row[0]).strip().lower() == "азимут":
            current["azimuth"] = _float(row[1])
        elif row and str(row[0]).strip().lower() == "высота":
            current["height"] = _float(row[1])
        for value in row[3:]:
            if value is not None:
                current["notes"].append(str(value))
    if current:
        result.append(SectionSoilBinding(**current))
    return result


def read_load_cases(path: str | Path) -> list[LoadCase]:
    if load_workbook is None:
        raise RuntimeError("openpyxl is required to read calculation workbooks")
    ws = load_workbook(path, data_only=True, read_only=True).active
    result: list[LoadCase] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(x is not None for x in row):
            continue
        result.append(
            LoadCase(
                number=str(row[0]),
                name=str(row[1]),
                load_type=str(row[2]),
                duration_share=_float(row[3]),
                reliability_factor=_float(row[4]),
            )
        )
    return result


def read_calculation_workbooks(
    bindings_path: str | Path,
    loads_path: str | Path,
) -> CalculationWorkbookProfile:
    return CalculationWorkbookProfile(
        bindings=read_section_soil_bindings(bindings_path),
        load_cases=read_load_cases(loads_path),
    )
