from docx import Document
from gip.analyzer import DocxReportAnalyzer
from gip.patch_planner import PatchPlanBuilder
from gip.patches import PatchKind

def test_analyzer_extracts_metadata_and_candidates(tmp_path):
    p=tmp_path/"r.docx"; d=Document()
    d.add_paragraph("Адрес: Ростовская область, г. Ростов-на-Дону")
    d.add_paragraph("Договор №123-45 от 01.09.2026")
    d.add_paragraph("Обнаружены трещины в наружной стене.")
    d.save(p)
    a=DocxReportAnalyzer().analyze(p)
    assert a.report.metadata.contract_number=="123-45"
    assert "01.09.2026" in a.report.metadata.dates
    assert a.report.metadata.address.startswith("Ростовская")
    assert len(a.candidate_defect_paragraphs)==1

def test_planner_requires_expected_values(tmp_path):
    p=tmp_path/"r.docx"; d=Document(); d.add_paragraph("Договор №123"); d.save(p)
    a=DocxReportAnalyzer().analyze(p)
    assert PatchPlanBuilder().build_metadata(a)==[]
    patches=PatchPlanBuilder().build_metadata(a,expected_contract="456")
    assert patches[0].kind is PatchKind.NORMAL
