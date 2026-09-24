from docx import Document
from gip.table_engine import DocxTableEngine

def make_doc(path):
    doc=Document()
    table=doc.add_table(rows=2, cols=3)
    table.cell(0,0).text="Дефект"
    table.cell(0,1).text="Категория состояния"
    table.cell(0,2).text="Примечание"
    table.cell(1,0).text="Трещина"
    table.cell(1,1).text="аварийное"
    table.cell(1,2).text="x"
    doc.save(path)

def test_reads_tables_and_rows(tmp_path):
    p=tmp_path/"r.docx"; make_doc(p)
    tables=DocxTableEngine().read_tables(p)
    assert tables[0].headers==["Дефект","Категория состояния","Примечание"]
    assert tables[0].rows[0][1]=="аварийное"

def test_category_validation(tmp_path):
    p=tmp_path/"r.docx"; make_doc(p)
    assert DocxTableEngine().validate_categories(p)==[]

def test_invalid_category_detected(tmp_path):
    p=tmp_path/"r.docx"; make_doc(p)
    doc=Document(p)
    doc.tables[0].cell(1,1).text="аварийное — требуется усиление"
    doc.save(p)
    result=DocxTableEngine().validate_categories(p)
    assert result[0].value=="аварийное — требуется усиление"
