from docx import Document
from docx.shared import RGBColor
from gip.docx_patch_engine import DocxPatchEngine
from gip.patches import Patch, PatchKind

def make(path):
    d=Document(); d.add_paragraph("Ошибка: старый текст"); d.save(path)

def test_normal_patch_is_green(tmp_path):
    p=tmp_path/"r.docx"; make(p)
    patch=Patch(kind=PatchKind.NORMAL, old="старый текст", new="исправленный текст", mark="green")
    DocxPatchEngine().apply(p,patch)
    r=Document(p).paragraphs[0].runs
    fixed=next(x for x in r if "исправленный" in x.text)
    assert fixed.font.color.rgb==RGBColor(0,128,0)

def test_date_patch_is_blue(tmp_path):
    p=tmp_path/"r.docx"; make(p)
    patch=Patch(kind=PatchKind.DATE, old="старый текст", new="новая дата", mark="blue")
    DocxPatchEngine().apply(p,patch)
    r=Document(p).paragraphs[0].runs
    fixed=next(x for x in r if "новая дата" in x.text)
    assert fixed.font.color.rgb==RGBColor(0,0,255)

def test_patch_can_find_table_cell(tmp_path):
    d=Document(); t=d.add_table(rows=1,cols=1); t.cell(0,0).text="старая категория"; d.save(tmp_path/"r.docx")
    patch=Patch(kind=PatchKind.NORMAL,old="старая категория",new="аварийное",mark="green")
    DocxPatchEngine().apply(tmp_path/"r.docx",patch)
    assert Document(tmp_path/"r.docx").tables[0].cell(0,0).text=="аварийное"
