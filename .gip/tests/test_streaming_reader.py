from docx import Document
from gip.streaming_reader import StreamingReportReader
from gip.analyzer import DocxReportAnalyzer


def test_reader_preserves_separate_tables_and_body(tmp_path):
    path = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("Разрушение штукатурного слоя стены.")
    t1 = doc.add_table(rows=2, cols=2)
    t1.cell(0, 0).text = "Номер"
    t1.cell(0, 1).text = "Дефект"
    t1.cell(1, 0).text = "ДКР-1"
    t1.cell(1, 1).text = "Трещина в стене"
    t2 = doc.add_table(rows=2, cols=2)
    t2.cell(0, 0).text = "Номер"
    t2.cell(0, 1).text = "Характеристика дефекта"
    t2.cell(1, 0).text = "ДКР-1"
    t2.cell(1, 1).text = "Вертикальная трещина"
    doc.save(path)

    parsed = StreamingReportReader().read(path)
    assert len(parsed.tables) == 2
    assert parsed.tables[0][1][0] == "ДКР-1"
    assert parsed.tables[1][1][0] == "ДКР-1"

    result = DocxReportAnalyzer().analyze(path)
    assert len(result.tables) == 2
    assert result.candidate_defect_paragraphs == ["Разрушение штукатурного слоя стены."]
