from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from gip.full_inspection import inspect_document

def make_doc(path):
    d=Document()
    d.add_paragraph("Адрес: Ростовская область, г. Ростов-на-Дону")
    d.add_paragraph("Договор №123")
    d.add_paragraph("Дата 01.09.2026")
    d.add_paragraph("Дата измерений 03.09.2026")
    d.add_paragraph("ДКР-1. Трещина в стене.")
    t=d.add_table(rows=3,cols=2)
    t.cell(0,0).text="№ дефекта"; t.cell(0,1).text="Категория состояния"
    t.cell(1,0).text="ДКР-1"; t.cell(1,1).text="аварийное"
    t.cell(2,0).text="ДКР-2"; t.cell(2,1).text="аварийное — пояснение"
    g=d.add_table(rows=2,cols=2)
    g.cell(0,0).text="Степень значимости"; g.cell(0,1).text="Описание"
    g.cell(1,0).text="Критический"; g.cell(1,1).text="Трещина"
    d.save(path)

def add_header_address(path):
    d=Document(path)
    header=d.sections[0].header
    p=header.paragraphs[0]
    p.text="Адрес: Ростовская область, г. Ростов-на-Дону, ул. Пушкинская, 10"
    d.save(path)

def test_full_checks_categories_and_graphics(tmp_path):
    p=tmp_path/"r.docx"; make_doc(p)
    codes={x.code for x in inspect_document(p).findings}
    assert "INVALID_CATEGORY" in codes

def test_significance_is_not_condition_category(tmp_path):
    p=tmp_path/"r.docx"; make_doc(p)
    codes={x.code for x in inspect_document(p).findings}
    assert "INVALID_CATEGORY" in codes
    # The "Степень значимости" table itself must not trigger category validation.

def test_multiple_dates_are_not_inconsistency(tmp_path):
    p=tmp_path/"r.docx"; make_doc(p)
    codes={x.code for x in inspect_document(p).findings}
    assert "DATE_INCONSISTENT" not in codes
    assert "DATE_MISMATCH" not in codes

def test_address_scans_header_and_keeps_special_semantics(tmp_path):
    p=tmp_path/"r.docx"; make_doc(p); add_header_address(p)
    x=next(x for x in inspect_document(p).findings if x.code=="ADDRESS_INCONSISTENT")
    assert x.patch_kind=="address"

def test_expected_date_is_explicit_only(tmp_path):
    p=tmp_path/"r.docx"; make_doc(p)
    x=next(x for x in inspect_document(p,expected_date="09.09.2026").findings if x.code=="DATE_MISMATCH")
    assert x.patch_kind=="date"
