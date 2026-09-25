from docx import Document

from gip.analyzer import DocxReportAnalyzer
from gip.large_docx import StreamingDocxReader


def test_streaming_reader_collects_body_and_header_without_whole_docx(tmp_path):
    path = tmp_path / "large-reader.docx"
    doc = Document()
    doc.add_paragraph("Договор №ОСК-ССК-22/0526-1")
    doc.add_paragraph("Дата обследования: 03.08.2026")
    doc.add_paragraph("Обнаружены трещины в наружной стене.")
    doc.sections[0].header.paragraphs[0].text = (
        "Адрес: Республика Крым, г. Симферополь, ул. Набережная, 28А"
    )
    doc.save(path)

    parsed = StreamingDocxReader().read(path)

    assert "word/document.xml" in parsed.paragraphs_by_part
    assert "word/header1.xml" in parsed.paragraphs_by_part
    assert parsed.text_chars > 0

    result = DocxReportAnalyzer().analyze(path)
    assert result.report.metadata.contract_number == "ОСК-ССК-22/0526-1"
    assert "03.08.2026" in result.report.metadata.dates
    assert any("Набережная" in value for value in result.address_candidates)
    assert len(result.candidate_defect_paragraphs) == 1
