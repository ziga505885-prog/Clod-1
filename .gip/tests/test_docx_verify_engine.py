from docx import Document
from docx.shared import RGBColor
from gip.docx_patch_engine import DocxPatchEngine
from gip.docx_verify_engine import DocxVerificationEngine
from gip.patches import Patch,PatchKind

def test_verify_green_patch(tmp_path):
    p=tmp_path/"r.docx"; d=Document(); d.add_paragraph("старый"); d.save(p)
    patch=Patch(PatchKind.NORMAL,"старый","новый","green")
    DocxPatchEngine().apply(p,patch)
    result=DocxVerificationEngine().verify_patch(p,patch)
    assert result.passed

def test_verify_blue_date(tmp_path):
    p=tmp_path/"r.docx"; d=Document(); d.add_paragraph("2025"); d.save(p)
    patch=Patch(PatchKind.DATE,"2025","2026","blue")
    DocxPatchEngine().apply(p,patch)
    assert DocxVerificationEngine().verify_patch(p,patch).passed

def test_verify_rejects_old_text(tmp_path):
    p=tmp_path/"r.docx"; d=Document(); d.add_paragraph("старый новый"); d.save(p)
    patch=Patch(PatchKind.NORMAL,"старый","новый","green")
    result=DocxVerificationEngine().verify_patch(p,patch)
    assert not result.passed
    assert any(x.code=="OLD_TEXT_PRESENT" for x in result.findings)

def test_verify_categories(tmp_path):
    p=tmp_path/"r.docx"; d=Document(); t=d.add_table(rows=2,cols=1); t.cell(0,0).text="Категория состояния"; t.cell(1,0).text="аварийное — пояснение"; d.save(p)
    result=DocxVerificationEngine().verify_document(p)
    assert not result.passed
