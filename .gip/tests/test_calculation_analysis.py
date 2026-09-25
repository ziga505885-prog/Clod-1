from pathlib import Path
from docx import Document
from gip.calculation_analysis import analyze_calculations

def test_calculation_analyzer_extracts_address_and_results(tmp_path: Path):
    p = tmp_path / "calc.docx"
    d = Document()
    d.add_paragraph("Комплекс: Бизнес-центр по адресу: Республика Крым, г. Симферополь, ул. Набережная, 28А")
    d.add_paragraph("Прогиб составил = 16 мм; Величина предельного прогиба: 30 мм.")
    d.add_paragraph("Фактическое армирование колонн превышает армирование, требуемое по расчету.")
    d.save(p)
    result = analyze_calculations([p], expected_address="Республика Крым, г. Симферополь, ул. Набережная, 28А")
    assert not any(f.code == "CALC_ADDRESS_MISMATCH" for f in result.findings)
    assert any("Прогиб составил" in x for x in result.results)
    assert result.standards == ()

def test_calculation_analyzer_does_not_invent_results(tmp_path: Path):
    p = tmp_path / "calc.docx"
    d = Document()
    d.add_paragraph("Расчетная схема.")
    d.save(p)
    result = analyze_calculations([p])
    assert any(f.code == "CALC_RESULTS_NOT_EXTRACTED" for f in result.findings)
