from docx import Document

from gip.large_docx import LargeDocxMetadataScanner


def test_large_docx_metadata_scanner_reads_headers_and_body(tmp_path):
    path = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("Договор №ОСК-ССК-22/0526-1")
    doc.add_paragraph("Дата обследования: 03.08.2026")
    doc.sections[0].header.paragraphs[0].text = (
        "Адрес: Республика Крым, г. Симферополь, ул. Набережная, 28А"
    )
    doc.save(path)

    result = LargeDocxMetadataScanner().scan(path)

    assert result.contracts == ["ОСК-ССК-22/0526-1"]
    assert "03.08.2026" in result.dates
    assert any("Набережная" in x for x in result.address_candidates)
    assert "word/document.xml" in result.part_names
    assert "word/header1.xml" in result.part_names
