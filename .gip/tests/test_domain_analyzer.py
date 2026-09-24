from docx import Document
from gip.analyzer import DocxReportAnalyzer
from gip.domain_analyzer import analyze_domain

def _make_doc(path):
    d=Document()
    d.add_paragraph("4 Характеристики дефектов")
    d.add_paragraph("ДКР-1. Трещина в наружной стене.")
    d.add_paragraph("Следы замачивания перекрытия.")
    t=d.add_table(rows=3, cols=2)
    t.cell(0,0).text="№ дефекта"; t.cell(0,1).text="Описание"
    t.cell(1,0).text="ДКР-1"; t.cell(1,1).text="Трещина в наружной стене"
    t.cell(2,0).text="ДКР-2"; t.cell(2,1).text="Коррозия металла"
    c=d.add_table(rows=2, cols=2)
    c.cell(0,0).text="№ дефекта"; c.cell(0,1).text="Характеристика дефектов"
    c.cell(1,0).text="ДКР-1"; c.cell(1,1).text="Трещина в наружной стене"
    d.save(path)

def test_domain_maps_summary_and_characteristics(tmp_path):
    p=tmp_path/"r.docx"; _make_doc(p)
    result=analyze_domain(DocxReportAnalyzer().analyze(p))
    assert {x.id for x in result.summary} == {"ДКР-1","ДКР-2"}
    assert {x.id for x in result.characteristics} == {"ДКР-1"}
    assert any(x.code=="DEFECT_MISSING_FROM_CHARACTERISTICS" for x in result.findings)

def test_domain_does_not_invent_id_for_general_defect(tmp_path):
    p=tmp_path/"r.docx"
    d=Document(); d.add_paragraph("Следы замачивания перекрытия."); d.save(p)
    result=analyze_domain(DocxReportAnalyzer().analyze(p))
    assert result.report[0].id is None

def test_domain_detects_duplicate_ids(tmp_path):
    p=tmp_path/"r.docx"
    d=Document()
    t=d.add_table(rows=3,cols=2)
    t.cell(0,0).text="№ дефекта"; t.cell(0,1).text="Описание"
    t.cell(1,0).text="ДКР-1"; t.cell(1,1).text="Трещина"
    t.cell(2,0).text="ДКР-1"; t.cell(2,1).text="Трещина повторно"
    d.save(p)
    result=analyze_domain(DocxReportAnalyzer().analyze(p))
    assert any(x.code=="DUPLICATE_DEFECT" for x in result.findings)
