from docx import Document
from gip.docx_patch_engine import DocxPatchEngine
from gip.docx_verify_engine import DocxVerificationEngine
from gip.patches import Patch, PatchKind

def test_patch_preserves_order(tmp_path):
    p=tmp_path/"r.docx"; d=Document(); d.add_paragraph("до старый после"); d.save(p)
    patch=Patch(PatchKind.NORMAL,"x","старый","новый","green","")
    DocxPatchEngine().apply(p,patch)
    assert Document(p).paragraphs[0].text=="до новый после"

def test_patch_spanning_runs(tmp_path):
    p=tmp_path/"r.docx"; d=Document(); para=d.add_paragraph()
    para.add_run("до ста"); para.add_run("рый после"); d.save(p)
    patch=Patch(PatchKind.NORMAL,"x","старый","новый","green","")
    DocxPatchEngine().apply(p,patch)
    assert Document(p).paragraphs[0].text=="до новый после"

def test_verify_uses_reason(tmp_path):
    p=tmp_path/"r.docx"; d=Document(); d.add_paragraph("2025"); d.save(p)
    patch=Patch(PatchKind.DATE,"x","2025","2026","blue","")
    DocxPatchEngine().apply(p,patch)
    assert DocxVerificationEngine().verify_patch(p,patch).passed

def test_verify_counts_occurrences(tmp_path):
    p=tmp_path/"r.docx"; d=Document(); d.add_paragraph("новый новый"); d.save(p)
    patch=Patch(PatchKind.NORMAL,"x","старый","новый","green","")
    result=DocxVerificationEngine().verify_patch(p,patch)
    assert not result.passed
    assert any(x.code=="REPLACEMENT_COUNT" for x in result.findings)
