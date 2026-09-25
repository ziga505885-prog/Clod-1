from docx import Document
from gip.analyzer import DocxReportAnalyzer

def test_gip_os_v03_header_address_and_multiple_dates(tmp_path):
    p = tmp_path / "metadata.docx"
    d = Document()
    d.add_paragraph("Дата обследования: 03.08.2026")
    d.add_paragraph("Дата отчета: 15.09.2026")
    d.add_paragraph("Сведения из архива: 23.06.2026")
    d.sections[0].header.paragraphs[0].text = (
        "Адрес: Республика Крым, г. Симферополь, ул. Набережная, 28А"
    )
    d.save(p)
    result = DocxReportAnalyzer().analyze(p)
    assert any("Набережная" in value for value in result.address_candidates)
    assert {"03.08.2026", "15.09.2026", "23.06.2026"} <= set(result.date_candidates)
