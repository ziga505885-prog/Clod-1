from docx import Document
from gip.pipeline import GIPPipeline
from gip.patches import Patch,PatchKind

def test_pipeline_full_cycle(tmp_path):
    source=tmp_path/"report.docx"
    d=Document(); d.add_paragraph("Старый текст"); d.save(source)
    target=tmp_path/"checked.docx"
    result=GIPPipeline().run(source,[Patch(PatchKind.NORMAL,"Старый текст","Новый текст","green")],target)
    assert result.completed
    assert result.applied_patches==1
    assert Document(target).paragraphs[0].text=="Новый текст"

def test_pipeline_rejects_bad_category(tmp_path):
    source=tmp_path/"report.docx"
    d=Document(); t=d.add_table(rows=2,cols=1); t.cell(0,0).text="Категория состояния"; t.cell(1,0).text="аварийное — пояснение"; d.save(source)
    result=GIPPipeline().run(source)
    assert not result.completed
    assert result.findings[0].code=="INVALID_CATEGORY"

def test_pipeline_copies_source_before_mutation(tmp_path):
    source=tmp_path/"report.docx"
    d=Document(); d.add_paragraph("старый"); d.save(source)
    result=GIPPipeline().run(source,[Patch(PatchKind.NORMAL,"старый","новый","green")])
    assert result.completed
    assert Document(source).paragraphs[0].text=="старый"
    assert result.output_path.exists()
